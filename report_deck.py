"""BPP (Budol-Proof Pinoy) brand awareness review deck, built from the same data layer as the Streamlit dashboard (data.py).

Output: a .pptx (native, editable charts and tables) in the Project Swarm look: white, navy block, yellow title word, red bar, Montserrat.
It opens in PowerPoint and Google Slides (import the file; Montserrat is a Google font, install it for PowerPoint).

Run (from this folder):
  python report_deck.py --month 2026-09               # data/bpp_snapshot.json (local, git-ignored)
  python report_deck.py --month 2026-09 --live        # read Supabase with .streamlit/secrets.toml or PULSE_DB_* variables
  python report_deck.py --month 2026-09 --internal    # adds the team-only data check (yellow months, Ads Manager vs Ad Spend Report)
  python report_deck.py --halves                      # half vs half review, 2 slides titled "Platform Distribution Level Up" (default 6 months from 2025-10)
  python report_deck.py --halves --layout halves      # same cards as first period on the left, next period on the right (cards | halves | table)
  python report_deck.py --halves --start 2025-10 --n 6
Without --month the last full month is used. Needs: pip install -r requirements-report.txt

The text is rule-based (no LLM). A person still checks the numbers against Meta and TikTok and writes the "why" and "what next"
(Media Planning Playbook 3.1, section 4.9): the closing slide leaves a box for the planner.
Awareness is reported as delivery and cost: reach measures exposure, not recall.
"""
import argparse
import os
import tomllib
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

import data as D

HERE = Path(__file__).parent
# Project Swarm palette (read from the deck file)
NAVY, BLUE, YELLOW, RED = "073763", "0B3691", "F1C232", "D10000"
TEXT, MUTED, LINE, MIST, WHITE = "3F3F3F", "737373", "DCE3EA", "F3F6FA", "FFFFFF"
GOOD, BAD = "3D8B37", "D10000"
FONT = "Montserrat"
W, H = 13.333, 7.5


# ---------- formatting ----------
def peso(x, dec=0):
    return "n/a" if pd.isna(x) else f"{'-' if x < 0 else ''}₱{abs(x):,.{dec}f}"


def num(x):
    return "n/a" if pd.isna(x) else f"{x:,.0f}"


def pct(x, dec=1):
    return "n/a" if pd.isna(x) else f"{x:.{dec}%}"


def short(x):
    return "n/a" if pd.isna(x) else (f"{x / 1e6:.1f} million" if x >= 1e6 else f"{x:,.0f}")


def mlabel(p):
    return p.strftime("%b %Y")


def change(c, p):
    return None if (c is None or p is None or pd.isna(c) or pd.isna(p) or p == 0) else c / p - 1


def tone(chg, higher_better, flat=0.02):
    """Colour of a change: green good, red bad, grey when flat or when the direction has no good or bad (higher_better None)."""
    if chg is None or pd.isna(chg) or abs(chg) < flat or higher_better is None:
        return MUTED
    return GOOD if (chg > 0) == higher_better else BAD


def rgb(h):
    return RGBColor.from_string(h)


# ---------- data ----------
def load_conf():
    """The same settings as the dashboard: .streamlit/secrets.toml [supabase], else PULSE_DB_* variables, else None (snapshot)."""
    f = HERE / ".streamlit" / "secrets.toml"
    if f.exists():
        conf = tomllib.loads(f.read_text()).get("supabase")
        if conf:
            return dict(conf)
    if os.environ.get("PULSE_DB_HOST"):
        return {"host": os.environ["PULSE_DB_HOST"], "user": os.environ["PULSE_DB_USER"], "password": os.environ["PULSE_DB_PASSWORD"],
                "port": os.environ.get("PULSE_DB_PORT", "5432"), "dbname": os.environ.get("PULSE_DB_NAME", "postgres"), "sslmode": os.environ.get("PULSE_DB_SSLMODE", "require")}
    return None


def build_report(S: dict, month: str | None, internal: bool) -> dict:
    ads, posts, tt = S["ads"], S["posts"], S["tiktok"]
    last = D.last_full_month(ads, S["built_at"])
    m = pd.Period(month, "M") if month else last
    p = m - 1
    months = sorted(ads["month"].unique())
    if m not in months:
        raise SystemExit(f"No Ads Manager data for {m}. Months available: {months[0]} to {months[-1]}.")
    has_prior = p >= months[0]
    cur = D.paid(ads, m, m)
    pre = D.paid(ads, p, p) if has_prior else None
    org, org_p = D.organic(posts, m, m), (D.organic(posts, p, p) if has_prior else None)
    tk, tk_p = D.tiktok(tt, m, m), (D.tiktok(tt, p, p) if has_prior else None)
    pm = D.paid_monthly(ads)
    pm = pm[pm.index <= m].tail(12)
    chk = D.spend_check(ads, S["ref"])
    gaps = set(chk[chk["difference"] < -1].index)
    camps = D.campaigns(ads, m, m)
    reach_camps = camps[(camps["reach"] > 0) & (camps["objective"] == "reach")]   # cost per 1,000 reached only compares reach campaigns
    R = {"S": S, "m": m, "label": mlabel(m), "vs": mlabel(p), "has_prior": has_prior, "internal": internal, "cur": cur, "pre": pre, "org": org, "org_p": org_p, "tk": tk,
         "tk_p": tk_p, "pm": pm, "gaps": gaps, "chk": chk, "camps": camps.head(8), "tp": D.top_posts(posts, m, m, 5), "bt": D.by_post_type(posts, m, m),
         "tv": D.top_tiktok(tt, m, m, 5), "partial": m >= pd.Period(S["built_at"], "M")}
    chg = lambda a, b, k: None if b is None else change(a[k], b[k])
    R["chg"] = {k: chg(cur, pre, k) for k in ("spend", "reach", "reach_avg", "frequency", "cpr", "thru_rate")}
    R["chg_org"] = {k: chg(org, org_p, k) for k in ("posts", "reach_per_post", "er", "views")}
    R["chg_tk"] = {k: chg(tk, tk_p, k) for k in ("videos", "views", "engagements", "eng_per_view")}
    # rule-based text, the same wording as the dashboard
    R["headline"] = (f"Facebook paid reach hit {short(cur['reach'])} in {R['label']}, at {peso(cur['cpr'], 2)} per 1,000 reached."
                     if not (pd.isna(cur["reach"]) or pd.isna(cur["cpr"])) else f"Brand awareness, {R['label']}")
    parts = []
    rc, cc = R["chg"]["reach"], R["chg"]["cpr"]
    if rc is not None and cc is not None:
        parts.append(f"That is {abs(rc):.1%} {'more' if rc >= 0 else 'less'} reach than {R['vs']}, at a cost per 1,000 {'up' if cc >= 0 else 'down'} {abs(cc):.1%}, "
                     f"with each person reached about {cur['frequency']:.1f} times on average.")
    if org["posts"]:
        parts.append(f"Facebook posts reached {num(org['reach_per_post'])} people each with a {pct(org['er'], 2)} engagement rate; TikTok videos drew {num(tk['views'])} views.")
    parts.append("Reach adds up campaign reach, so the number of different people is lower.")
    R["standfirst"] = " ".join(parts)
    ask = "Agree target audiences per platform so reach can be reported as a share of audience and judged against a target."
    if len(reach_camps) >= 2:
        lo, hi = reach_camps.loc[reach_camps["cpr"].idxmin()], reach_camps.loc[reach_camps["cpr"].idxmax()]
        R["focus"] = (f"Among reach campaigns, {lo['name'][:60]} was cheapest at {peso(lo['cpr'], 2)} per 1,000 reached and {hi['name'][:60]} the costliest at "
                      f"{peso(hi['cpr'], 2)}. {ask}")
    else:
        R["focus"] = ask
    return R


def build_halves(S: dict, start: str, n: int, internal: bool, layout: str = "cards") -> dict:
    """Half vs half review: the first n months from `start` against the next n months, with the cards of the dashboard's KPI Cards tab (data.half_cards)."""
    ads, posts, tt = S["ads"], S["posts"], S["tiktok"]
    a0 = pd.Period(start, "M")
    last = D.last_full_month(ads, S["built_at"])
    if a0 + (2 * n - 1) > last:
        raise SystemExit(f"{n} + {n} months from {mlabel(a0)} run past the last full month ({mlabel(last)}).")
    HC = D.half_cards(ads, posts, tt, a0, n, S["acct"], S["follows"], S["tt_platform"])
    for g in HC["groups"]:     # the follows per ₱1,000 card has no monthly series in data.py; the key card needs one for its sparkline
        for row in g[1]:
            for c in row:
                if c["label"] == "New follows per ₱1,000" and len(S["follows"]):
                    fm = dict(zip(S["follows"]["month"], S["follows"]["follows"].astype(float)))
                    c["series"] = [D._div(fm[m] * 1000, D.paid(ads, m, m)["spend"]) for m in HC["months"]]
    (a0, a1), (b0, b1) = HC["a"], HC["b"]
    rng = lambda x, y: f"{mlabel(x)} to {mlabel(y)}"
    short_rng = lambda x, y: f"{x.strftime('%b')}–{mlabel(y)}" if x.year == y.year else f"{mlabel(x)}–{mlabel(y)}"   # Apr–Sep 2026, Oct 2025–Mar 2026
    R = {"S": S, "HC": HC, "n": n, "page": S["page"], "internal": internal, "layout": layout, "a_rng": rng(a0, a1), "b_rng": rng(b0, b1), "a_short": short_rng(a0, a1), "b_short": short_rng(b0, b1),
         "vs": f"first {n} months",
         "label": f"{rng(a0, a1)} vs {rng(b0, b1)}", "file": f"BPP_Awareness_Halves_{a0}_{n}m_{layout}"}
    return R


# ---------- deck ----------
class Deck:
    font = FONT

    def __init__(self, R):
        self.R = R
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.n = 0

    # --- primitives ---
    def rect(self, s, x, y, w, h, fill, line=None, lw=1.0):
        b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        b.fill.solid(); b.fill.fore_color.rgb = rgb(fill)
        if line:
            b.line.color.rgb = rgb(line); b.line.width = Pt(lw)
        else:
            b.line.fill.background()
        b.shadow.inherit = False
        st = b._element.find(qn("p:style"))   # drop the theme style reference: it brings a drop shadow in some viewers
        if st is not None:
            b._element.remove(st)
        return b

    def text(self, s, text, x, y, w, h, size, bold=False, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        p.alignment = align
        self.run(p, text, size, bold, color)
        return tb

    def run(self, p, text, size, bold=False, color=TEXT):
        r = p.add_run()
        r.text = text
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, self.font
        r.font.color.rgb = rgb(color)
        return r

    def new(self):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(WHITE)
        return s

    def slide(self, lead, accent, sub=None):
        """Content slide: navy bar top-left, red bar bottom-right, title with the last words in yellow, footer."""
        s = self.new()
        self.rect(s, 0, 0, W * 0.38, 0.14, NAVY)
        self.rect(s, W * 0.5, H - 0.14, W * 0.5, 0.14, RED)
        tb = self.text(s, f"{lead} ", 0.6, 0.45, 12.1, 0.7, 28, color=NAVY)
        self.run(tb.text_frame.paragraphs[0], accent, 28, True, YELLOW)
        if sub:
            self.text(s, sub, 0.6, 1.17, 12.1, 0.3, 12, color=MUTED)
        R = self.R
        self.text(s, f"BPP · Brand awareness review · {R['label']} · Data as of {R['S']['built_at']:%d %b %Y}", 0.6, 6.88, 9.5, 0.25, 9, color=MUTED)
        self.text(s, str(self.n), 11.9, 6.88, 0.8, 0.25, 9, color=MUTED, align=PP_ALIGN.RIGHT)
        return s

    def note(self, s, text, y=6.4, h=0.4):
        self.text(s, text, 0.6, y, 12.1, h, 9.5, color=MUTED)

    def kpi(self, s, x, y, w, label, value, line, chg, higher_better, h=1.6):
        self.rect(s, x, y, w, h, MIST)
        self.rect(s, x, y, w, 0.06, NAVY)
        self.text(s, label, x + 0.18, y + 0.2, w - 0.3, 0.3, 10.5, bold=True, color=MUTED)
        self.text(s, value, x + 0.18, y + 0.52, w - 0.3, 0.5, 24, bold=True, color=NAVY)
        self.text(s, line, x + 0.18, y + 1.05, w - 0.3, 0.2, 9, color=MUTED)
        t = "No prior month to compare" if chg is None else f"{chg:+.1%} vs {self.R['vs']}"
        self.text(s, t, x + 0.18, y + 1.3, w - 0.3, 0.2, 10, bold=True, color=tone(chg, higher_better))

    def takeaway(self, s, x, y, w, h, blocks, size=13):
        self.rect(s, x, y, w, h, WHITE, line=NAVY, lw=1.25)
        self.rect(s, x, y, 0.12, h, YELLOW)
        tb = s.shapes.add_textbox(Inches(x + 0.35), Inches(y + 0.2), Inches(w - 0.6), Inches(h - 0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        first = True
        for k, v in blocks:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            self.run(p, k, 11, True, BLUE)
            p2 = tf.add_paragraph(); p2.space_after = Pt(8)
            self.run(p2, v, size, False, TEXT)

    @staticmethod
    def _borders(cell):
        """Thin bottom line only (the default table style draws white grid lines). Border elements must come before the fill in tcPr."""
        from lxml import etree

        ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
        tcPr = cell._tc.get_or_add_tcPr()
        for tag in ("lnB", "lnT", "lnR", "lnL"):
            ln = etree.SubElement(tcPr, f"{{{ns}}}{tag}")
            tcPr.remove(ln)
            if tag == "lnB":
                ln.set("w", "9525")
                etree.SubElement(etree.SubElement(ln, f"{{{ns}}}solidFill"), f"{{{ns}}}srgbClr").set("val", LINE)
            else:
                ln.set("w", "0")
                etree.SubElement(ln, f"{{{ns}}}noFill")
            tcPr.insert(0, ln)

    def table(self, s, x, y, headers, rows, widths, aligns, row_h=0.44, size=10.5):
        gt = s.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(sum(widths)), Inches(row_h * (len(rows) + 1)))
        tbl = gt.table
        tblPr = gt._element.graphic.graphicData.tbl.tblPr
        for a in ("bandRow", "firstRow"):
            tblPr.set(a, "0")
        for j, wd in enumerate(widths):
            tbl.columns[j].width = Inches(wd)
        for i in range(len(rows) + 1):
            tbl.rows[i].height = Inches(row_h)
            for j in range(len(headers)):
                c = tbl.cell(i, j)
                c.fill.solid(); c.fill.fore_color.rgb = rgb(NAVY if i == 0 else WHITE)
                c.margin_left = c.margin_right = Inches(0.1)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                self._borders(c)
                p = c.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT if aligns[j] == "r" else PP_ALIGN.LEFT
                self.run(p, str(headers[j] if i == 0 else rows[i - 1][j]), size, bold=(i == 0 or j == 0), color=WHITE if i == 0 else (NAVY if j == 0 else TEXT))

    def line_chart(self, s, x, y, w, h, cats, vals, numfmt, label_fmt, flags=None):
        vals = list(vals)
        cd = CategoryChartData(); cd.categories = cats
        cd.add_series("Monthly value", [None if pd.isna(v) else float(v) for v in vals])
        flagged = [i for i, f in enumerate(flags or []) if f]
        if flagged:   # internal view: months whose Ads Manager export is probably incomplete, as their own marker-only series (renders in every viewer)
            cd.add_series("Export probably incomplete", [None if (i not in flagged or pd.isna(vals[i])) else float(vals[i]) for i in range(len(vals))])
        ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
        ch.has_title = False
        ch.has_legend = False
        ch.font.size, ch.font.name = Pt(9), FONT
        ch.font.color.rgb = rgb(MUTED)
        va = ch.value_axis
        va.major_gridlines.format.line.color.rgb = rgb("E9EEF3")
        va.format.line.fill.background()
        va.tick_labels.number_format, va.tick_labels.number_format_is_linked = numfmt, False
        ch.category_axis.format.line.color.rgb = rgb(LINE)
        ch.category_axis.has_major_gridlines = False
        ser = ch.plots[0].series[0]
        ser.smooth = False
        ser.format.line.color.rgb = rgb(BLUE); ser.format.line.width = Pt(2.5)
        ser.marker.style, ser.marker.size = XL_MARKER_STYLE.CIRCLE, 7
        ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(BLUE); ser.marker.format.line.color.rgb = rgb(NAVY)
        if flagged:
            s2 = ch.plots[0].series[1]
            s2.smooth = False
            s2.format.line.fill.background()
            s2.marker.style, s2.marker.size = XL_MARKER_STYLE.CIRCLE, 10
            s2.marker.format.fill.solid(); s2.marker.format.fill.fore_color.rgb = rgb(YELLOW); s2.marker.format.line.color.rgb = rgb(NAVY)
        valid = [i for i, v in enumerate(vals) if not pd.isna(v)]
        if valid:
            dl = ser.points[valid[-1]].data_label
            dl.position = XL_LABEL_POSITION.ABOVE
            tf = dl.text_frame
            tf.text = label_fmt(vals[valid[-1]])
            tf.paragraphs[0].runs[0].font.size, tf.paragraphs[0].runs[0].font.bold = Pt(10), True
            tf.paragraphs[0].runs[0].font.color.rgb = rgb(NAVY)
        return ch

    # --- slides ---
    def build(self):
        R, c, ch, org, tk, pm = self.R, self.R["cur"], self.R["chg"], self.R["org"], self.R["tk"], self.R["pm"]

        # 1. cover (the Swarm cover: navy block top-left, red bar bottom-right, right-aligned title)
        s = self.new()
        self.rect(s, 0, 0, 5.1, 3.7, NAVY)
        self.rect(s, 6.8, H - 0.74, W - 6.8, 0.74, RED)
        self.rect(s, 0, H - 0.74, 5.1, 0.74, RED)
        self.text(s, "BUDOL-PROOF PINOY", 5.6, 2.2, 7.1, 0.7, 30, color=BLUE, align=PP_ALIGN.RIGHT)
        self.text(s, "Brand Awareness", 4.6, 2.95, 8.1, 1.1, 48, bold=True, color=YELLOW, align=PP_ALIGN.RIGHT)
        self.text(s, "MONTHLY REVIEW", 5.6, 4.55, 7.1, 0.4, 18, bold=True, color=TEXT, align=PP_ALIGN.RIGHT)
        self.text(s, R["label"], 5.6, 4.95, 7.1, 0.4, 18, bold=True, color=TEXT, align=PP_ALIGN.RIGHT)
        self.text(s, "Draft: numbers to be checked against Meta and TikTok before this goes to the client.", 5.6, 5.55, 7.1, 0.4, 10, color=MUTED, align=PP_ALIGN.RIGHT)

        # 2. headline
        s = self.slide("Brand awareness,", R["label"], "Paid awareness on Facebook, with the Facebook page posts (including boosted) and TikTok alongside")
        self.text(s, R["headline"], 0.6, 1.75, 11.6, 1.3, 30, bold=True, color=NAVY)
        self.text(s, R["standfirst"], 0.6, 3.2, 11.4, 1.1, 14, color=TEXT)
        heroes = [("Reach", num(c["reach"]), ch["reach"], True), ("Cost per 1,000 reached", peso(c["cpr"], 2), ch["cpr"], False), ("Average frequency", f"{c['frequency']:.2f}", ch["frequency"], False)]
        for i, (lab, val, chg_, hb) in enumerate(heroes):
            x = 0.6 + i * 4.1
            self.rect(s, x, 4.6, 3.8, 0.05, NAVY)
            self.text(s, lab, x, 4.75, 3.8, 0.3, 11, bold=True, color=MUTED)
            self.text(s, val, x, 5.05, 3.8, 0.7, 36, bold=True, color=NAVY)
            self.text(s, "No prior month to compare" if chg_ is None else f"{chg_:+.1%} vs {R['vs']}", x, 5.8, 3.8, 0.3, 11, bold=True, color=tone(chg_, hb))
        self.note(s, "Reach is the sum of campaign reach in the Ads Manager export, and campaigns overlap, so unique reach is lower. Reach measures exposure, not recall.", y=6.4)

        # 3. paid awareness
        s = self.slide("Paid", "awareness", f"Facebook, Ads Manager, {R['label']}. Targets are pending sign-off, so nothing here is judged on target yet.")
        items = [("Ad spend", peso(c["spend"]), f"{c['campaigns']} campaigns", ch["spend"], None), ("Reach", num(c["reach"]), "Share of target audience: pending", ch["reach"], True),
                 ("Average frequency", f"{c['frequency']:.2f}", "Monthly; alert above 3 is weekly", ch["frequency"], False), ("Cost per 1,000 reached", peso(c["cpr"], 2), "Target pending", ch["cpr"], False),
                 ("Video view-through", pct(c["thru_rate"], 1), "ThruPlays / impressions", ch["thru_rate"], True)]
        for i, (lab, val, line, chg_, hb) in enumerate(items):
            self.kpi(s, 0.6 + i * 2.45, 1.75, 2.3, lab, val, line, chg_, hb)
        self.takeaway(s, 0.6, 3.75, 12.1, 2.35, [("Key takeaway", R["standfirst"]), ("Suggested focus (auto-generated)", R["focus"])])
        self.note(s, "Cost per 1,000 reached only compares reach campaigns: engagement boosts optimise for engagement, not reach.")

        # 4. awareness trend
        s = self.slide("Awareness", "Trend", f"Paid awareness on Facebook by month, up to {R['label']}. Weekly reach and target audiences are not available yet, so months are shown.")
        cats = [m.strftime("%b %y") for m in pm.index]
        flags = [m in R["gaps"] for m in pm.index] if R["internal"] else None
        specs = [("Reach (sum of campaign reach)", pm["reach"], "#,##0", short), ("Average frequency (impressions / reach)", pm["frequency"], "0.0", lambda v: f"{v:.2f}"),
                 ("Cost per 1,000 reached (₱)", pm["cpr"], "₱#,##0", lambda v: peso(v, 2)), ("Video view-through (ThruPlays / impressions)", pm["thru_rate"], "0%", lambda v: pct(v, 1))]
        for i, (title, vals, fmt, lab) in enumerate(specs):
            x, y = 0.6 + (i % 2) * 6.15, 1.65 + (i // 2) * 2.4
            self.text(s, title, x, y, 5.9, 0.25, 11, bold=True, color=NAVY)
            self.line_chart(s, x - 0.1, y + 0.25, 6.0, 2.05, cats, vals, fmt, lab, flags)
        n = "Reach is the sum of campaign reach per month, so it overstates unique reach. Frequency is monthly; the deck's alert level of 3 applies to weekly frequency."
        if R["internal"] and R["gaps"]:
            n = "INTERNAL: yellow markers are months whose Ads Manager export has less spend than the Ad Spend Report, so the export is probably incomplete. " + n
        self.note(s, n, y=6.4)

        # 5. paid campaigns
        s = self.slide("Paid", "campaigns", f"{R['label']}, largest spend first")
        rows = [[r["name"][:62], r["objective"], peso(r["spend"], 2), num(r["reach"]), f"{r['frequency']:.2f}" if not pd.isna(r["frequency"]) else "n/a", peso(r["cpr"], 2), pct(r["thru_rate"], 1)]
                for _, r in R["camps"].iterrows()]
        self.table(s, 0.6, 1.7, ["Campaign", "Type", "Spend", "Reach", "Freq.", "Cost / 1,000", "View-through"], rows, [4.7, 1.2, 1.35, 1.3, 0.9, 1.4, 1.25], ["l", "l", "r", "r", "r", "r", "r"], size=9.5)
        self.note(s, "Boosting campaigns optimise for engagement, so their cost per 1,000 reached is not comparable with reach campaigns.", y=6.1)

        # 6. facebook page
        s = self.slide("Facebook", "page", f"Page posts published in {R['label']}, including boosted posts (lifetime numbers of each post as exported)")
        co = R["chg_org"]
        for i, (lab, val, line, chg_) in enumerate([("Posts published", num(org["posts"]), "Facebook page", co["posts"]), ("Average reach per post", num(org["reach_per_post"]), "Lifetime reach", co["reach_per_post"]),
                                                    ("Engagement rate", pct(org["er"], 2), "Engagements / reach", co["er"]), ("Post views", num(org["views"]), "Facebook post views", co["views"])]):
            self.kpi(s, 0.6 + i * 3.05, 1.7, 2.9, lab, val, line, chg_, True, h=1.55)
        tp = R["tp"]
        rows = [[" ".join(str(r["title"]).split())[:70], r["type"], f"{r['published']:%d %b}", num(r["views"]), num(r["reach"]), pct(r["er"], 1)] for _, r in tp.iterrows()]
        self.text(s, "Top posts by reach", 0.6, 3.5, 6, 0.3, 13, bold=True, color=NAVY)
        if rows:
            self.table(s, 0.6, 3.9, ["Post", "Type", "Date", "Views", "Reach", "Eng. rate"], rows, [6.0, 1.3, 1.0, 1.3, 1.3, 1.2], ["l", "l", "l", "r", "r", "r"], row_h=0.4, size=9.5)

        # 7. tiktok
        s = self.slide("TikTok", "videos", f"Videos posted in {R['label']}. TikTok has views and engagement only in the source data, so there is no reach.")
        ct = R["chg_tk"]
        for i, (lab, val, line, chg_) in enumerate([("Videos posted", num(tk["videos"]), "TikTok", ct["videos"]), ("Video views", num(tk["views"]), f"{num(tk['views_per_video'])} per video", ct["views"]),
                                                    ("Engagements", num(tk["engagements"]), "Likes, comments, shares", ct["engagements"]), ("Engagement per view", pct(tk["eng_per_view"], 1), "Engagements / views", ct["eng_per_view"])]):
            self.kpi(s, 0.6 + i * 3.05, 1.7, 2.9, lab, val, line, chg_, True, h=1.55)
        tv = R["tv"]
        rows = [[" ".join(str(r["title"]).split())[:70], f"{r['posted']:%d %b}", num(r["views"]), num(r["likes"]), num(r["shares"]), pct(r["eng_per_view"], 1)] for _, r in tv.iterrows()]
        self.text(s, "Top videos by views", 0.6, 3.5, 6, 0.3, 13, bold=True, color=NAVY)
        if rows:
            self.table(s, 0.6, 3.9, ["Video", "Date", "Views", "Likes", "Shares", "Eng. / view"], rows, [6.0, 1.0, 1.3, 1.3, 1.3, 1.2], ["l", "l", "r", "r", "r", "r"], row_h=0.4, size=9.5)

        # 8. next steps and notes
        s = self.slide("What to look", "at next", "Written by the dashboard from the numbers; the planner adds the why and the plan")
        self.takeaway(s, 0.6, 1.75, 6.2, 1.9, [("Suggested focus (auto-generated)", R["focus"])])
        self.rect(s, 0.6, 3.9, 6.2, 2.3, WHITE, line=LINE, lw=1.25)
        self.text(s, "Planner's notes", 0.85, 4.05, 5.7, 0.3, 11, bold=True, color=BLUE)
        self.text(s, "Why it happened, and what we do next. To be written by the planner before this goes to the client.", 0.85, 4.4, 5.7, 1.0, 12, color=MUTED)
        self.text(s, "About these numbers", 7.3, 1.75, 5.4, 0.3, 13, bold=True, color=NAVY)
        tb = s.shapes.add_textbox(Inches(7.3), Inches(2.15), Inches(5.4), Inches(4.0))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        bullets = ["Reach is the sum of the campaign reach rows of the monthly Ads Manager export. Campaigns overlap, so unique reach is lower.",
                   "Frequency is impressions / reach over the month. The deck's alert level of 3 is weekly.",
                   "Cost per 1,000 reached = ad spend / reach x 1,000.",
                   "Engagement rate = reactions + comments + shares, divided by reach, for Facebook page posts.",
                   "Not available yet: weekly reach, target audiences (share of audience), follower growth, brand health.",
                   "Awareness is reported as delivery and cost. Reach measures exposure, not recall."]
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(7)
            self.run(p, "• " + b, 11, False, TEXT)
        if R["partial"]:
            self.note(s, f"{R['label']} was still running when the data was loaded: its numbers are partial.", y=6.4)
        return self.prs


class HalvesDeck(Deck):
    """Half vs half review in the look of the "Share-Led Content Focus" deck: cream page, dark ink, deep green for results that improved and burnt orange for
    results that got worse, DM Sans, white rounded cards and a dark note bar. Two slides under one title, "Platform Distribution Level Up".
    The cards come from data.half_cards, the same list the dashboard's KPI Cards tab uses. Each card reads in the order Sir Joven asked for: the number now,
    a filled change pill, then the previous number and the difference."""
    font = "DM Sans"
    CREAM, INK, TEAL, ORANGE, ASH, EDGE, SAGE = "F6F4EF", "1B2A34", "17695D", "B4480F", "4F5D66", "D9D5CA", "A9B8B3"
    TITLE = "Platform Distribution Level Up"
    X0, CW = 0.9, 11.55
    KIND = {"num": (num, 0), "num1": (lambda v: "n/a" if pd.isna(v) else f"{v:,.1f}", 1), "peso0": (peso, 0), "peso2": (lambda v: peso(v, 2), 2), "peso3": (lambda v: peso(v, 3), 3),
            "rate": (lambda v: pct(v, 2), 4), "freq": (lambda v: "n/a" if pd.isna(v) else f"{v:.2f}", 2)}
    # big cards (with a monthly sparkline) and table rows of each slide. "big" lists candidates in order: the first three that exist become the key cards and the
    # rest join the rows, so the deck still builds when follows are not loaded. A card that is not listed is added to the table of the second slide, so nothing
    # silently drops out
    LAYOUT = [{"big": ["Average reach per month", "Facebook followers", "Average reach per Facebook post", "Facebook post views (impressions)", "New Facebook follows", "TikTok video views"],
               "rows": ["TikTok followers", "TikTok reach", "TikTok engagements", "Total reach", "Ad impressions", "Unique people reached per month", "Reach per ₱1,000",
                        "Impressions per ₱1,000", "TikTok engagement per view"]},
              {"big": ["New follows per ₱1,000", "Engagements per ₱1,000", "Cost per engagement", "Cost per mille (CPM)"],
               "rows": ["Ad post engagements", "Ad engagement rate", "Click-through rate (CTR)", "Cost per click", "Cost per 1,000 reached", "Ad spend per month", "Average frequency",
                        "True frequency", "Facebook posts published", "TikTok videos posted", "Facebook engagement rate"]}]
    # Some campaign- or video-based cards have a more accurate twin, so the deck shows only the twin when both exist: true frequency in place of the average frequency, and the TikTok analytics engagement in place of the engagement per view of the videos.
    # The dashboard still shows both. Empty this dict to bring the campaign-based cards back.
    SUPERSEDED = {"Average frequency": ("True frequency",), "TikTok engagement per view": ("TikTok engagements",)}
    NAMES = {"Facebook post views (impressions)": "Post views (impressions)", "Average reach per Facebook post": "Reach per Facebook post", "Unique people reached per month": "Unique reach per month"}   # the longest title wraps to two lines in a wider font; the speaker notes say it is an average

    def name(self, c):
        return self.NAMES.get(c["label"], c["label"])

    TAGS = {"Average reach per month": "FACEBOOK ADS", "Facebook followers": "FACEBOOK PAGE", "New Facebook follows": "FACEBOOK PAGE", "Average reach per Facebook post": "FACEBOOK PAGE", "Facebook post views (impressions)": "FACEBOOK PAGE", "TikTok video views": "TIKTOK",
            "New follows per ₱1,000": "FACEBOOK PAGE + ADS", "Engagements per ₱1,000": "FACEBOOK ADS", "Cost per engagement": "FACEBOOK ADS", "Cost per mille (CPM)": "FACEBOOK ADS"}

    # --- primitives ---
    def box(self, s, x, y, w, h, fill, line=None, adj=0.04):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        b.adjustments[0] = adj
        b.fill.solid(); b.fill.fore_color.rgb = rgb(fill)
        if line:
            b.line.color.rgb = rgb(line); b.line.width = Pt(1)
        else:
            b.line.fill.background()
        b.shadow.inherit = False
        st_ = b._element.find(qn("p:style"))   # drop the theme style reference: it brings a drop shadow in some viewers
        if st_ is not None:
            b._element.remove(st_)
        return b

    def tone(self, c):
        """Colour of a card: green when it got better, orange when worse, ash when about the same (within 2%), not compared, or neither better nor worse."""
        chg = change(c["b"], c["a"]) if c["compared"] else None
        if chg is None or c["better"] is None or abs(chg) < 0.02:
            return self.ASH
        return self.TEAL if (chg > 0) == c["better"] else self.ORANGE

    @staticmethod
    def pill_text(c):
        chg = change(c["b"], c["a"]) if c["compared"] else None
        return "Not compared" if not c["compared"] else ("n/a" if chg is None else f"{'▲' if chg > 0 else '▼'} {abs(chg):.1%}")

    def pill(self, s, x, y, w, h, c, size):
        t = self.pill_text(c)
        p = self.box(s, x, y, w, h, self.tone(c), adj=0.22)
        tf = p.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        self.run(tf.paragraphs[0], t, size, True, WHITE)

    def diff(self, c):
        """Absolute change as text: '+1,200', '−₱0.50', '+0.40 pts'; empty when either period has no value."""
        if not c["compared"] or pd.isna(c["a"]) or pd.isna(c["b"]):
            return ""
        fmt, rnd = self.KIND[c["kind"]]
        d = round(c["b"], rnd) - round(c["a"], rnd)
        return ("+" if d >= 0 else "−") + (f"{abs(d) * 100:.2f} pts" if c["kind"] == "rate" else fmt(abs(d)))

    def sparkline(self, s, x, y, w, h, vals):
        """Monthly values of both periods, drawn as plain shapes (a line per period and a dot per month) instead of a chart: Google Slides shows the axes of an
        imported chart even when they are hidden, and shapes look the same in PowerPoint, Google Slides and Keynote. The second line starts at the first
        period's last month so the two join."""
        n = len(vals) // 2
        pts = [(i, float(v)) for i, v in enumerate(vals) if v is not None and not pd.isna(v)]
        if len(pts) < 2:
            return
        lo, hi = min(v for _, v in pts), max(v for _, v in pts)
        span = (hi - lo) or 1.0
        pad, r = 0.07, 0.035
        grp = s.shapes.add_group_shape()    # one object per sparkline, so it moves as a piece when edited
        px = lambda i: x + pad + (w - 2 * pad) * i / (len(vals) - 1)
        py = lambda v: y + pad + (h - 2 * pad) * (1 - (v - lo) / span)
        for a, b, colour in ((0, n - 1, self.SAGE), (n - 1, len(vals) - 1, self.TEAL)):
            seg = [(i, v) for i, v in pts if a <= i <= b]
            if len(seg) < 2:
                continue
            fb = grp.shapes.build_freeform(Inches(px(seg[0][0])), Inches(py(seg[0][1])))
            fb.add_line_segments([(Inches(px(i)), Inches(py(v))) for i, v in seg[1:]], close=False)
            line = fb.convert_to_shape()
            line.fill.background()
            line.line.color.rgb = rgb(colour); line.line.width = Pt(2)
            line.shadow.inherit = False
            st_ = line._element.find(qn("p:style"))
            if st_ is not None:
                line._element.remove(st_)
        for i, v in pts:
            dot = grp.shapes.add_shape(MSO_SHAPE.OVAL, Inches(px(i) - r), Inches(py(v) - r), Inches(2 * r), Inches(2 * r))
            dot.fill.solid(); dot.fill.fore_color.rgb = rgb(self.SAGE if i < n else self.TEAL)
            dot.line.fill.background()
            dot.shadow.inherit = False
            st_ = dot._element.find(qn("p:style"))
            if st_ is not None:
                dot._element.remove(st_)

    # --- page pieces ---
    def page(self, subtitle, source):
        s = self.new()
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(self.CREAM)
        self.text(s, self.TITLE, self.X0, 0.5, self.CW, 0.6, 28, True, self.INK)
        self.text(s, subtitle, self.X0, 1.12, self.CW, 0.55, 13, color=self.ASH)
        long = len(source) > 520
        vlong = len(source) > 700
        self.text(s, source, self.X0, 6.78 if vlong else (6.8 if long else 6.88), self.CW, 0.65 if vlong else (0.55 if long else 0.45), 8.5 if vlong else (9 if long else 10), color=self.ASH)
        return s

    def big_card(self, s, x, y, w, h, c, compact=False):
        """A key KPI: platform tag, name, the number now, the change pill, the previous number with the difference, and the monthly sparkline.
        compact=True is for slides where a second row of cards sits below (a slightly smaller number and tighter spacing)."""
        self.box(s, x, y, w, h, WHITE, self.EDGE)
        fmt, _ = self.KIND[c["kind"]]
        pad, n = 0.22, self.R["n"]
        ny, size, py = (0.78, 30, 1.34) if compact else (0.84, 34, 1.5)
        if c["label"] in self.TAGS:
            self.text(s, self.TAGS[c["label"]], x + pad, y + 0.14, w - 2 * pad, 0.2, 10.5, True, self.TEAL)
        self.text(s, self.name(c), x + pad, y + 0.34 if compact else y + 0.37, w - 2 * pad, 0.5, 15, True, self.INK)
        t = self.tone(c)
        self.text(s, fmt(c["b"]), x + pad, y + ny, w - 2 * pad, 0.6, size, True, self.INK if t == self.ASH else t)
        py += y
        pw = 0.092 * len(self.pill_text(c)) + 0.3
        self.pill(s, x + pad, py, pw, 0.3, c, 12)
        tb = self.text(s, "vs ", x + pad + pw + 0.12, py, w - 2 * pad - pw - 0.12, 0.3, 13, False, self.INK, anchor=MSO_ANCHOR.MIDDLE)
        self.run(tb.text_frame.paragraphs[0], fmt(c["a"]), 13, True, self.INK)
        if self.diff(c):
            self.text(s, f"{self.diff(c)} vs the first {n} months", x + pad, py + (0.34 if compact else 0.36), w - 2 * pad, 0.22, 11, color=self.ASH)
        if c["series"] is not None:
            self.sparkline(s, x + 0.1, y + h - 0.47, w - 0.2, 0.42, c["series"])

    def small_card(self, s, x, y, w, h, c, tight=False):
        """A secondary KPI in the same reading order as the big cards: name, the number now, the change pill with the previous number.
        tight=True is for five cards in a row (smaller type), used when more cards than usual have to share the slide."""
        self.box(s, x, y, w, h, WHITE, self.EDGE)
        fmt, _ = self.KIND[c["kind"]]
        six = tight == 6
        pad = 0.14 if six else (0.18 if tight else 0.2)
        k = len(self.name(c))
        if six:
            size = 9.5
        elif tight:
            size = 10.5 if k <= 18 else 9.5
        else:
            size = 12 if k <= 20 else (11 if k <= 24 else 10.5)
        self.text(s, self.name(c), x + pad, y + 0.12, w - 2 * pad, 0.24, size, True, self.INK)
        t = self.tone(c)
        self.text(s, fmt(c["b"]), x + pad, y + 0.36, w - 2 * pad, 0.45, 20 if six else (22 if tight else 24), True, self.INK if t == self.ASH else t)
        py = y + 0.86
        pw = (0.07 if six else 0.085) * len(self.pill_text(c)) + (0.2 if six else 0.28)
        self.pill(s, x + pad, py, pw, 0.27, c, 9 if six else (10.5 if tight else 11))
        vsz = 9 if six else (11 if tight else 12)
        tb = self.text(s, "vs ", x + pad + pw + 0.1, py, w - 2 * pad - pw - 0.1, 0.27, vsz, False, self.INK, anchor=MSO_ANCHOR.MIDDLE)
        self.run(tb.text_frame.paragraphs[0], fmt(c["a"]), vsz, True, self.INK)

    def half_panel(self, s, x, y, w, row_h, tag, rng, cards, second, stack):
        """One half of the half vs half layout, like the dashboard's first tab: a panel for the first period (the numbers with a Baseline pill) and one for the next
        period (the numbers with the change pill), metrics in the same grid cell in both so each pair sits level."""
        cols = 3 if stack else 2     # stacked cells (pill under the number) are narrower, so a third column keeps long numbers and names clear of each other
        h = 0.8 + -(-len(cards) // cols) * row_h + 0.12
        self.box(s, x, y, w, h, "E9F2EF" if second else WHITE, self.TEAL if second else self.EDGE)
        self.text(s, tag, x + 0.25, y + 0.15, w - 0.5, 0.22, 11, True, self.TEAL)
        self.text(s, rng, x + 0.25, y + 0.36, w - 0.5, 0.28, 14, False, self.ASH)
        cw = (w - 0.5) / cols
        for i, c in enumerate(cards):
            r, col = divmod(i, cols)
            cx, cy = x + 0.25 + col * cw, y + 0.8 + r * row_h
            fmt, _ = self.KIND[c["kind"]]
            self.text(s, self.name(c), cx, cy, cw - 0.12, 0.4 if stack else 0.22, 10.5 if stack else 11, False, self.ASH)
            t = self.tone(c)
            size, ny = (22, 0.38) if stack else (20, 0.2)
            self.text(s, fmt(c["b"] if second else c["a"]), cx, cy + ny, cw - 0.1, 0.45, size, True, t if (second and t != self.ASH) else self.INK)
            label = self.pill_text(c) if second else "Baseline"
            pw = 0.085 * len(label) + 0.28
            px, py = (cx, cy + 0.84) if stack else (cx + cw - 0.12 - pw, cy + 0.27)
            if second:
                self.pill(s, px, py, pw, 0.26, c, 10.5)
            else:
                p = self.box(s, px, py, pw, 0.26, "EAE6DC", adj=0.22)
                tf = p.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf.paragraphs[0].alignment = PP_ALIGN.CENTER
                self.run(tf.paragraphs[0], label, 10.5, True, self.ASH)
        return y + h

    def table(self, s, y, cards, max_h=2.0):
        """The remaining KPIs as rows: name, the number now, the change pill, the previous number and the difference."""
        fmt_of = lambda c: self.KIND[c["kind"]][0]
        cols = {"name": 0.9, "val": 5.2, "chg": 7.0, "prev": 8.7, "diff": 10.5}
        for k, label in (("name", "Metric"), ("val", f"Next {self.R['n']} months"), ("chg", "Change"), ("prev", f"First {self.R['n']} months"), ("diff", "Difference")):
            self.text(s, label, cols[k], y, 1.8 if k != "name" else 4, 0.26, 11, True, self.ASH)
        rh = min(0.32, max_h / max(len(cards), 1))
        y0 = y + 0.3
        for i, c in enumerate(cards):
            ry = y0 + i * rh
            self.rect(s, self.X0, ry, self.CW, 0.01, self.EDGE)
            fmt = fmt_of(c)
            t = self.tone(c)
            mid = MSO_ANCHOR.MIDDLE
            self.text(s, self.name(c), cols["name"], ry, 4.2, rh, 13, True, self.INK, anchor=mid)
            self.text(s, fmt(c["b"]), cols["val"], ry, 1.7, rh, 14, True, self.INK if t == self.ASH else t, anchor=mid)
            ph = min(0.24, rh - 0.04)
            self.pill(s, cols["chg"], ry + (rh - ph) / 2, 1.2, ph, c, 11)
            self.text(s, fmt(c["a"]), cols["prev"], ry, 1.7, rh, 13, False, self.INK, anchor=mid)
            self.text(s, self.diff(c), cols["diff"], ry, 1.9, rh, 12, False, self.ASH, anchor=mid)
        self.rect(s, self.X0, y0 + len(cards) * rh, self.CW, 0.01, self.EDGE)
        return y0 + len(cards) * rh

    def bar(self, s, y, label, text):
        """The dark note bar of the reference deck: a bold lead-in word and one sentence."""
        self.box(s, self.X0, y, self.CW, 0.58, self.INK, adj=0.25)
        tb = self.text(s, f"{label} ", self.X0 + 0.25, y, self.CW - 0.5, 0.58, 12, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
        self.run(tb.text_frame.paragraphs[0], text, 12, False, WHITE)

    @staticmethod
    def notes(s, text):
        s.notes_slide.notes_text_frame.text = text

    def spec(self):
        """What the two slides say and which cards they hold; the layouts below only decide how the cards are drawn."""
        R, HC, n = self.R, self.R["HC"], self.R["n"]
        by = {c["label"]: c for g in HC["groups"] for row in g[1] for c in row}
        listed = {l for sl in self.LAYOUT for k in ("big", "rows") for l in sl[k]}
        extras = [c for l, c in by.items() if l not in listed]       # cards the layout does not know yet (click-through rate, cost per click)
        mv = lambda c: f"{'up' if change(c['b'], c['a']) >= 0 else 'down'} {abs(change(c['b'], c['a'])):.1%}"
        vs = f"({R['b_short']} vs {R['a_short']})"
        up = lambda c: c["compared"] and (change(c["b"], c["a"]) or 0) > 0

        def pick(i):
            """The first three candidates are the key cards; the other candidates and the listed rows follow as small cards or table rows."""
            cand = [by[l] for l in self.LAYOUT[i]["big"] if l in by]
            rows = [by[l] for l in self.LAYOUT[i]["rows"] if l in by and not any(t in by for t in self.SUPERSEDED.get(l, ()))]
            return cand[:3], cand[3:] + rows

        # slide 1: how far the content travelled and how many people followed
        big, rows = pick(0)
        fol, rpp, ttv = by.get("Facebook followers") or by.get("New Facebook follows"), by.get("Average reach per Facebook post"), by.get("TikTok video views")
        lead = [c for c in (fol, rpp, ttv) if c and up(c)]
        names = {"Facebook followers": "Facebook followers", "New Facebook follows": "new Facebook follows", "Average reach per Facebook post": "reach per Facebook post", "TikTok video views": "TikTok video views"}
        avg = by.get("Average reach per month")
        if avg and avg["compared"] and not pd.isna(avg["b"]) and change(avg["b"], avg["a"]) is not None:
            _c = change(avg["b"], avg["a"])
            tail = (f", with average ad reach holding at {num(avg['b'])} a month" if abs(_c) < 0.02 else f", with average ad reach per month {mv(avg)} to {num(avg['b'])}")
        else:
            tail = ""
        if len(lead) >= 2:
            sub = f"The audience and the content's reach grew: " + ", ".join(f"{names[c['label']]} {mv(c)}" for c in lead[:-1]) + f" and {names[lead[-1]['label']]} {mv(lead[-1])}{tail} {vs}."
        else:
            sub = f"How far content travelled on Facebook and TikTok{tail} {vs}."
        posts, vids = by.get("Facebook posts published"), by.get("TikTok videos posted")
        counts = (f"{int(posts['a'] + posts['b'])} Facebook posts and {int(vids['a'] + vids['b'])} TikTok videos, " if posts and vids else "")
        foot = (f"Source: BPP Content SSoT 2025–2026, TikTok Studio and Meta Business Suite; {counts}{R['a_rng']} and {R['b_rng']}. Post views are Meta's Views of page posts: every time a post "
                "was shown or played, repeats included, organic and boosted together (lifetime figures of the posts published in each period). Unique reach is Meta's account-level monthly figure; "
                "total reach sums the monthly campaign reach and ad impressions count every time an ad was shown, so people reached more than once count each time.")
        plat = any(l in by for l in ("Facebook followers", "TikTok followers", "TikTok reach", "TikTok engagements"))
        if plat:
            foot += " Followers, TikTok reach and TikTok engagements are the platforms' own figures for each period, not sums of posts or videos."
        pg = R.get("page") or {}
        whole = ((f"Over the whole period the page gained {num(pg['net_follows'])} net followers. " if "Facebook followers" in by else
                  f"Over the whole period the page gained {num(pg['net_follows'])} net followers and has {num(pg['followers_lifetime'])} in total. ")
                 if pg.get("net_follows") and pg.get("followers_lifetime") else "")
        bar = ("Note:", (f"New follows are counted before unfollows. {whole}" if fol else "") + "Facebook page posts include boosted posts, so part of this reach is paid.")
        notes = (f"How to read each card: the big number is the latest {n} months ({R['b_rng']}); the filled pill is the change against the first {n} months ({R['a_rng']}), "
                 "green when it got better, orange when it got worse, grey when it is about the same (within 2%) or neither better nor worse; the line underneath is the previous number. "
                 f"The small chart, where there is one, shows each month: the first {n} months in light green and the next {n} in dark green.\n\n"
                 "Facebook followers, TikTok followers, TikTok reach and TikTok engagements are the platforms' own figures for each period, supplied by the team (they are not sums of posts or videos, and TikTok reach is not the same as TikTok video views). "
                 "New Facebook follows is the total of the daily new follows in Meta Business Suite for the months of the period. It is counted before unfollows: unfollows are only known for the whole year, "
                 "so net new followers per period cannot be shown. Total reach adds up the monthly campaign reach of the Facebook ads, so people reached in several months or by several campaigns are counted each time. "
                 "Unique reach per month is Meta's own count of different people the ads reached in a month (account-level report), averaged over the period. Reach per Facebook post is the average reach of the posts published in the period (lifetime numbers as exported); Facebook page posts include boosted posts. "
                 "Post views (impressions) are Meta's Views of the page posts published in the period: every time a post was shown or played, repeats included, organic and boosted posts together, so they are not a number of people and not purely organic. "
                 "Ad impressions are every time a Facebook ad was shown (Meta's account totals). TikTok video views are totals for the period. Reach per ₱1,000 still uses campaign reach from Ads Manager, which adds up campaigns, so the number of different people is lower. "
                 "Reach and impressions per ₱1,000 show how much distribution each ₱1,000 of ad spend bought. TikTok engagement per view is likes, comments and shares divided by views.")
        table_h = 1.4
        if (self.R["layout"] != "halves" and len(rows) > 5) or (self.R["layout"] == "halves" and len(big) + len(rows) > 12):    # no room for the note bar under that many cards: a short form joins the footnote
            foot += (" New follows are counted before unfollows" + (f"; the page gained {num(pg['net_follows'])} net followers over the whole period" if pg.get("net_follows") else "")
                     + ". Facebook page posts include boosted posts, so part of this reach is paid.")
            bar, table_h = None, 1.95
        out = [dict(big=big, rows=rows, sub=sub, foot=foot, bar=bar, notes=notes, stack=True, table_h=table_h)]

        # slide 2: what each peso delivered, and how much we posted
        big, rows = pick(1)
        rows[0:0] = extras
        fpk, eng, cpe = by.get("New follows per ₱1,000"), by.get("Engagements per ₱1,000"), by.get("Cost per engagement")
        if eng and cpe and change(eng["b"], eng["a"]) > 0 and change(cpe["b"], cpe["a"]) < 0:
            parts = ([f"new follows per ₱1,000 {mv(fpk)}"] if fpk and up(fpk) else []) + [f"engagements per ₱1,000 {mv(eng)}"]
            sub = f"Each peso of ad spend delivered more: {', '.join(parts)} and cost per engagement {mv(cpe)} {vs}."
        else:
            sub = f"What each peso of ad spend delivered, and how much we posted {vs}."
        worse = [c for c in rows if c["better"] is True and c["compared"] and (change(c["b"], c["a"]) or 0) < -0.02]
        if worse:   # the orange figures, said once in words so the headline does not read as all good news
            sub += f" {len(worse)} of the figures below went down."
        foot = ("Source: Facebook Ads Manager monthly exports, ad spend ex-VAT; click-through rate, cost per click and true frequency use Meta's account-level monthly totals. "
                "Reach per ₱1,000 and cost per 1,000 reached use campaign reach, which adds up people who saw several campaigns. New follows per ₱1,000 is all new page follows "
                "(Meta Business Suite, before unfollows) divided by Facebook ad spend, so it includes follows that did not come from ads. Ad post engagements are the post engagements of the Facebook ads as counted by Ads Manager. "
                "The difference of a rate is in percentage points (pts).")
        if R["internal"]:
            foot += (" INTERNAL: the daily follows total is about 4% above what Meta's whole-period audience figures imply (unexplained); treat new follows as approximate.")
        notes = ("Same reading as the previous slide. New follows per ₱1,000, engagements per ₱1,000 and impressions per ₱1,000 are what each ₱1,000 of Facebook ad spend bought (higher is better). "
                 "New follows per ₱1,000 divides all new page follows, organic and paid, by Facebook ad spend, so it is not the cost of a follower from ads. "
                 "Ad post engagements is the total number of post engagements of the Facebook ads as counted by Ads Manager. "
                 "Cost per engagement, cost per mille (CPM, cost per 1,000 impressions) and cost per 1,000 reached are what each result cost (lower is better). "
                 "Ad engagement rate is post engagements divided by impressions. Click-through rate is clicks divided by impressions and cost per click is spend divided by clicks, both from Meta's account-level monthly totals. "
                 "True frequency is impressions divided by unique reach. Ad spend per month and true frequency are context, not good or bad, so they stay grey. "
                 "Facebook posts published and TikTok videos posted show how much we posted; the Facebook engagement rate is engagements divided by reach of the page posts.")
        watch = ("Down:", ", ".join(c["label"] for c in worse) + ".") if worse else None
        out.append(dict(big=big, rows=rows, sub=sub, foot=foot, bar=None, watch=watch, notes=notes, stack=False, table_h=1.95))
        return out

    # --- layouts ---
    def draw_table(self, s, sp):
        """Three key cards, then the other figures as table rows."""
        cw = (self.CW - 0.46) / 3
        bh = 2.4 if sp["bar"] else 2.62      # a slide with a note bar has less room: smaller key cards keep the bar clear of the source line
        for i, c in enumerate(sp["big"]):
            self.big_card(s, self.X0 + i * (cw + 0.23), 1.75, cw, bh, c, compact=bool(sp["bar"]))
        y = self.table(s, 1.75 + bh + 0.12 if sp["bar"] else 4.52, sp["rows"], max_h=sp["table_h"])
        if sp["bar"]:
            self.bar(s, y + 0.2, *sp["bar"])

    def draw_cards(self, s, sp):
        """Everything as cards: three key cards with monthly charts, then rows of smaller cards in the same reading order."""
        cw = (self.CW - 0.46) / 3
        for i, c in enumerate(sp["big"]):
            self.big_card(s, self.X0 + i * (cw + 0.23), 1.72, cw, 2.4, c, compact=True)
        rows = sp["rows"]
        cols = 6 if len(rows) > 10 else (5 if (len(rows) > 8 or len(rows) == 5) else 4)    # five (six when there are more than ten) to a row when a slide holds many small cards, so the rows stay even
        sw, y = (self.CW - 0.12 * (cols - 1)) / cols, 4.26
        for k in range(0, len(rows), cols):
            for i, c in enumerate(rows[k:k + cols]):
                self.small_card(s, self.X0 + i * (sw + 0.12), y, sw, 1.2, c, tight=6 if cols == 6 else cols == 5)
            y += 1.3
        if sp.get("watch") and len(rows) % cols:   # the free slot at the end of the last row: a dark card in the style of the reference deck
            i = len(rows) % cols
            fx, fw = self.X0 + i * (sw + 0.12), sw * (cols - i) + 0.12 * (cols - i - 1)
            self.box(s, fx, y - 1.3, fw, 1.2, self.INK, adj=0.08)
            tb = self.text(s, sp["watch"][0] + " ", fx + 0.25, y - 1.3, fw - 0.5, 1.2, 13, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
            self.run(tb.text_frame.paragraphs[0], sp["watch"][1], 13, False, WHITE)
        if sp["bar"]:
            self.bar(s, y + 0.05, *sp["bar"])

    def draw_halves(self, s, sp):
        """The half vs half format of the dashboard's first tab: the first period on the left, the next period on the right, every metric level across both."""
        n, R = self.R["n"], self.R
        cards = sp["big"] + sp["rows"]
        pw = (self.CW - 0.3) / 2
        stack = sp["stack"] and len(cards) <= 9        # more than nine cards no longer fit three to a row with the pill under the number: use two columns with the pill beside it
        rows_n = -(-len(cards) // (3 if stack else 2))
        row_h = 1.15 if stack else min(0.78, (6.65 - 1.75 - 0.92 - (0.7 if sp["bar"] else 0)) / rows_n)   # more metrics squeeze the rows; a note bar needs its own room below
        y = self.half_panel(s, self.X0, 1.75, pw, row_h, f"FIRST {n} MONTHS", R["a_rng"], cards, False, stack)
        self.half_panel(s, self.X0 + pw + 0.3, 1.75, pw, row_h, f"NEXT {n} MONTHS", R["b_rng"], cards, True, stack)
        if sp["bar"]:
            self.bar(s, y + 0.12, *sp["bar"])

    def build(self):
        draw = {"table": self.draw_table, "cards": self.draw_cards, "halves": self.draw_halves}[self.R["layout"]]
        for sp in self.spec():
            s = self.page(sp["sub"], sp["foot"])
            draw(s, sp)
            self.notes(s, sp["notes"])
        return self.prs


def main():
    ap = argparse.ArgumentParser(description="Build the BPP brand awareness review deck")
    ap.add_argument("--month", help="YYYY-MM (default: last full month)")
    ap.add_argument("--live", action="store_true", help="read Supabase (secrets.toml or PULSE_DB_* variables) instead of the local snapshot")
    ap.add_argument("--internal", action="store_true", help="add the team-only data check (yellow months, note)")
    ap.add_argument("--out", help="output file (default out/BPP_Awareness_YYYY-MM.pptx)")
    ap.add_argument("--halves", action="store_true", help="half vs half review (the dashboard's KPI cards) instead of the monthly review")
    ap.add_argument("--start", default="2025-10", help="with --halves: first month of the first period (default 2025-10)")
    ap.add_argument("--n", type=int, default=6, help="with --halves: months per period (default 6)")
    ap.add_argument("--layout", choices=["cards", "halves", "table"], default="cards",
                    help="with --halves: cards (every figure a card), halves (first period left, next period right, like the dashboard's first tab) or table (key cards plus table rows)")
    a = ap.parse_args()
    conf = load_conf() if a.live else None
    if a.live and not conf:
        raise SystemExit("--live needs a [supabase] section in .streamlit/secrets.toml or PULSE_DB_* environment variables.")
    if not a.live and not D.snapshot_exists():
        raise SystemExit("No local snapshot (data/bpp_snapshot.json). Use --live to read Supabase.")
    S = D.load(conf)
    if a.halves:
        R = build_halves(S, a.start, a.n, a.internal, a.layout)
        out = Path(a.out) if a.out else HERE / "out" / f"{R['file']}.pptx"
        out.parent.mkdir(parents=True, exist_ok=True)
        HalvesDeck(R).build().save(out)
        print(f"Saved {out}  ({R['label']}, source: {S['source']})")
        return
    R = build_report(S, a.month, a.internal)
    out = Path(a.out) if a.out else HERE / "out" / f"BPP_Awareness_{R['m']}.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    Deck(R).build().save(out)
    print(f"Saved {out}  ({R['label']}, source: {S['source']})")


if __name__ == "__main__":
    main()
